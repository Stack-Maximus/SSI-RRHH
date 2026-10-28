# -*- coding: utf-8 -*-
"""
Plantilla en blanco, en Excel, de un comprobante tal como lo emite SSI-RRHH-SST
(mismo diseño y mismos campos que el PDF "Concept 3" del sistema).

- Las piezas gráficas (encabezado, barras de sección, línea del pie e íconos)
  son imágenes dibujadas con el mismo motor del sistema (sgc-layout.js), SIN
  los textos variables (ver assets.cjs).
- Encima van cuadros de texto editables, vinculados a celdas: título,
  código, fecha y revisión del encabezado; número y nombre de cada sección.
- Las fichas y la tabla de aprobaciones son celdas, con la geometría del PDF.

    node assets.cjs <tipo> 0.6     # imágenes + doc.json
    python3 generar.py <tipo>      # out/<código>_Comprobante_<Tipo>.xlsx
"""
import os, re, sys, json, zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.fills import GradientFill, Stop
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils import get_column_letter as L

BASE = os.path.dirname(os.path.abspath(__file__))
EMU = 9525
FUENTE = 'Arial'

PAL = dict(fondo='F7F9FB', blanco='FFFFFF', borde='E2EAF2', navy='021659', texto='3A5A85', azulEtiqueta='0A64B4',
           ficha='F1F6FA', separador='D4DFE9', filaAlterna='F0F5F9', lineaTabla='E3EBF2', bordeTabla='DCE5EE',
           pieFuerte='00417D', pieTexto='557AA2', pieSep='9AB0C8', numero='1A396C')
DEGRADADO = [(0, '003174'), (0.72, '00357B'), (0.87, '0058A6'), (1, '0A7FC3')]
FICHA = dict(x=[50.5, 723], w=645, wFull=1317.5, h=88, gap=15.6, iconCx=47.5, iconSize=44, sepX=94)
TABLA = dict(x=38.5, w=1342, hHead=87, hMin=88)
FIJOS = {'Registro', 'Código', 'Revisión'}
LARGOS = re.compile(r'^(motivo|modificaciones|observaciones|hito de la obra|nueva condición)$', re.I)


def fill(hex_):
    return PatternFill('solid', fgColor='FF' + hex_)


def sin_relleno(cel):
    f = cel.fill
    return f is None or (getattr(f, 'fill_type', None) is None and not hasattr(f, 'stop'))


def color_en(t):
    t = max(0, min(1, t))
    for (o0, c0), (o1, c1) in zip(DEGRADADO, DEGRADADO[1:]):
        if t <= o1:
            f = (t - o0) / ((o1 - o0) or 1)
            return ''.join('%02X' % round(int(c0[j:j + 2], 16) * (1 - f) + int(c1[j:j + 2], 16) * f) for j in (0, 2, 4))
    return DEGRADADO[-1][1]


class Comprobante:
    def __init__(self, tipo):
        self.tipo = tipo
        self.dir = os.path.join(BASE, 'assets', tipo)
        meta = json.load(open(os.path.join(self.dir, 'doc.json')))
        self.K = meta['K']
        self.doc, self.textos, self.imgs = meta['doc'], meta['textos'], meta['imagenes']
        self.cuadros = []
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = 'Maestro' if tipo == 'maestro' else 'Comprobante'

        # ---- grilla de columnas: une los cortes de encabezado, tarjetas, fichas y tabla (en u) ----
        X0 = 26  # la columna B empieza donde empieza el encabezado
        cortes = {26, 33, 1384.5, 1392}
        cortes |= {50.5, 50.5 + 94, 695.5, 723, 723 + 94, 1368}
        tx = TABLA['x']
        acc = tx
        anchos_t = self._anchos_tabla()
        for w in anchos_t:
            cortes.add(round(acc, 3)); acc += w
        cortes.add(round(acc, 3))
        # bordes extra: donde empiezan los cuadros de texto y los íconos (así
        # quedan anclados en un borde de columna, sin desplazamiento)
        cortes |= {709, 67.5 - 25, 152.5, FICHA['x'][0] + FICHA['iconCx'] - FICHA['iconSize'] / 2,
                   FICHA['x'][1] + FICHA['iconCx'] - FICHA['iconSize'] / 2}
        for t in self.textos['encabezado']:
            if t['f'] in ('mBold', 'f700'):
                cortes.add(t['x'])
        sec_t = next(s for s in self.doc['secciones'] if s.get('tabla'))
        acc = tx
        for c, w in zip(sec_t['tabla']['columnas'], anchos_t):
            if c.get('icono'):
                cortes.add(round(acc + 17, 3))
            acc += w
        self.cortes = sorted(round(u, 3) for u in cortes)
        self.px_corte = {u: round((u - X0) * self.K) for u in self.cortes}
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 16 / 7
        for i in range(len(self.cortes) - 1):
            px = self.px_corte[self.cortes[i + 1]] - self.px_corte[self.cortes[i]]
            ws.column_dimensions[L(2 + i)].width = px / 7
        self.col_fin = 2 + len(self.cortes) - 1  # columna de margen derecho
        ws.column_dimensions[L(self.col_fin)].width = 16 / 7
        self.fila = 1
        self.y_px = 0  # alto acumulado (px) desde la fila 1

    def _anchos_tabla(self):
        sec = next(s for s in self.doc['secciones'] if s.get('tabla'))
        cols = sec['tabla']['columnas']
        tot = sum(c['w'] for c in cols)
        return [c['w'] / tot * TABLA['w'] for c in cols]

    # columna (1-based) donde empieza el corte u
    def col(self, u):
        return 2 + self.cortes.index(round(u, 3) if round(u, 3) in self.cortes else u)

    def x_px(self, u):  # px desde el inicio de la columna B
        return (u - 26) * self.K

    def ancla_u(self, x_u):
        if round(x_u, 3) in self.cortes:
            return self.col(x_u) - 1, 0
        return self.ancla(self.x_px(x_u))

    def ancla(self, x_px_desde_B):
        # (columna 0-based, desplazamiento px) para un x en px desde el inicio de la columna B
        for i in range(len(self.cortes) - 1):
            a, b = self.px_corte[self.cortes[i]], self.px_corte[self.cortes[i + 1]]
            if x_px_desde_B < b:
                return 1 + i, x_px_desde_B - a
        return self.col_fin - 2, x_px_desde_B - self.px_corte[self.cortes[-2]]

    def nueva_fila(self, alto_u=None, alto_px=None):
        px = alto_px if alto_px is not None else max(1, round(alto_u * self.K))
        self.ws.row_dimensions[self.fila].height = px * 0.75
        r = self.fila
        self.fila += 1
        self.y_px += px
        return r

    def unir(self, r0, c0, r1, c1):
        if (r0, c0) != (r1, c1):
            self.ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)
        return self.ws.cell(r0, c0)

    def pintar(self, r0, c0, r1, c1, hex_):
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                self.ws.cell(r, c).fill = fill(hex_)

    def borde_caja(self, r0, c0, r1, c1, hex_):
        s = Side(style='thin', color='FF' + hex_)
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                cel = self.ws.cell(r, c)
                b = cel.border
                cel.border = Border(left=s if c == c0 else b.left, right=s if c == c1 else b.right,
                                    top=s if r == r0 else b.top, bottom=s if r == r1 else b.bottom)

    def imagen(self, nombre, fila, x_u, dy_px=0, ancho=None, alto=None):
        ruta = os.path.join(self.dir, nombre + '.png')
        meta = self.imgs[nombre]
        w, h = ancho or meta['ancho'], alto or meta['alto']
        img = XLImage(ruta)
        col, off = self.ancla_u(x_u)
        img.anchor = OneCellAnchor(_from=AnchorMarker(col=col, colOff=int(off * EMU), row=fila - 1, rowOff=int(dy_px * EMU)),
                                   ext=XDRPositiveSize2D(int(w * EMU), int(h * EMU)))
        self.ws.add_image(img)

    def imagen_estirada(self, nombre, fila0, fila1, x0=26, x1=1392):
        """Imagen que ocupa exactamente de la fila0 a la fila1 (inclusive) y del borde x0 al x1:
        se estira con las celdas, así calza igual en Excel y en otros programas."""
        img = XLImage(os.path.join(self.dir, nombre + '.png'))
        img.anchor = TwoCellAnchor(editAs='oneCell',
                                   _from=AnchorMarker(col=self.col(x0) - 1, colOff=0, row=fila0 - 1, rowOff=0),
                                   to=AnchorMarker(col=self.col(x1) - 1, colOff=0, row=fila1, rowOff=0))
        self.ws.add_image(img)

    def cuadro(self, nombre, fila, fila_y0_px, x_u, y_u_centro, w_u, s_u, texto, color, b=True, algn='l', vinculo=None,
               fila_ref_y_px=None, wrap=True, alto_u=None):
        """Cuadro de texto. y_u_centro: centro vertical en u, relativo al tope de `fila` (fila_y0_px = px del tope)."""
        K = self.K
        h_u = alto_u or s_u * 1.45
        col, off = self.ancla_u(x_u)
        top_px = y_u_centro * K - h_u * K / 2
        self.cuadros.append(dict(nombre=nombre, col=col, colOff=int(off * EMU), row=fila - 1, rowOff=int(max(0, top_px) * EMU),
                                 cx=int(w_u * K * EMU), cy=int(h_u * K * EMU), texto=texto,
                                 sz=int(round(s_u * K * 0.75 * 100 / 10) * 10), color=color, b=b, algn=algn,
                                 vinculo=vinculo, wrap=wrap))

    # ------------------------------------------------------------ partes
    def encabezado(self):
        enc = self.doc['encabezado']
        K = self.K
        alto_img = self.imgs['encabezado']['alto']
        # 4 filas bajo la imagen (en la 2 queda el título, fuente de su cuadro de texto)
        partes = [alto_img // 4] * 3 + [alto_img - 3 * (alto_img // 4)]
        filas = [self.nueva_fila(alto_px=p) for p in partes]
        self.pintar(filas[0], 1, filas[-1], self.col_fin, PAL['fondo'])
        cel_tit = self.ws.cell(filas[1], 2)
        cel_tit.value = '\n'.join(x['txt'] for x in self.textos['encabezado'] if x['f'] == 'mBold')
        cel_tit.font = Font(name=FUENTE, size=10, bold=True, color='FF' + PAL['navy'])
        self.imagen_estirada('encabezado', filas[0], filas[-1])
        t = self.textos['encabezado']
        tit = [x for x in t if x['f'] == 'mBold']
        y_top = tit[0]['y'] - tit[0]['s'] * 0.95 - 34
        y_bot = tit[-1]['y'] + tit[-1]['s'] * 0.32 - 34
        ancho_titulo = 984.9 - (166 - 45) * 0.2232 - 22 - tit[0]['x']  # el mismo ancho que usa el sistema
        self.cuadro('TituloSGC', filas[0], 0, tit[0]['x'], (y_top + y_bot) / 2, ancho_titulo,
                    tit[0]['s'], '\n'.join(x['txt'] for x in tit), 'FFFFFF', vinculo=f'$B${filas[1]}',
                    alto_u=(y_bot - y_top) + 10)
        self.enc_valores = [x for x in t if x['f'] == 'f700']  # código, fecha, revisión (en ese orden)
        self.filas_enc = filas
        self.nueva_fila(alto_u=10)  # 322 -> 332

    def cuadros_encabezado(self, refs):
        nombres = ['ValorCodigo', 'ValorFecha', 'ValorRevision']
        f0 = self.filas_enc[0]
        for nombre, val, ref in zip(nombres, self.enc_valores, refs):
            self.cuadro(nombre, f0, 0, val['x'], val['y'] - 34 - val['s'] * 0.36, 1376 - val['x'], val['s'],
                        val['txt'], val['c'].lstrip('#'), vinculo=ref, wrap=False)

    def barra(self, n, titulo):
        r = self.nueva_fila(alto_u=59.5)
        self.pintar(r, 1, r, self.col_fin, PAL['fondo'])
        c_tit = self.ws.cell(r, 2)
        c_tit.value = titulo.upper()
        c_tit.font = Font(name=FUENTE, size=10, bold=True, color='FF' + PAL['fondo'])
        self.imagen_estirada(f'barra{n}', r, r)
        tx = {x['f']: x for x in self.textos[f'barra{n}']}
        num, tit = tx['mBlack'], tx['mBold']
        self.cuadro(f'NumSeccion{n}', r, 0, num['x'] - 25, num['y'] - num['s'] * 0.36, 50, num['s'], num['txt'],
                    num['c'].lstrip('#'), algn='ctr')
        self.cuadro(f'Seccion{n}', r, 0, tit['x'], tit['y'] - tit['s'] * 0.36, 900, tit['s'], titulo.upper(), 'FFFFFF',
                    vinculo=f'$B${r}')

    def fichas(self, fichas):
        K = self.K
        c_card0, c_card1 = self.col(33), self.col(1384.5) - 1
        r_ini = self.nueva_fila(alto_u=16.5)
        # filas de fichas: de a dos, las largas a lo ancho
        filas, i = [], 0
        while i < len(fichas):
            a = fichas[i]
            if a.get('ancho') == 'completo':
                filas.append([('completo', a)]); i += 1
            elif i + 1 < len(fichas) and fichas[i + 1].get('ancho') != 'completo':
                filas.append([('izq', a), ('der', fichas[i + 1])]); i += 2
            else:
                filas.append([('izq', a)]); i += 1
        celdas = {}
        self.filas_dobles = []
        for k, fila in enumerate(filas):
            if k:
                rg = self.nueva_fila(alto_u=FICHA['gap'])
                if len(fila) == 2 and len(filas[k - 1]) == 2:
                    self.filas_dobles.append(rg)
            largo = any(p == 'completo' and LARGOS.match(c['etiqueta']) for p, c in fila)
            r_et = self.nueva_fila(alto_u=34)
            r_val = self.nueva_fila(alto_u=54 + (60 if largo else 0))
            if len(fila) == 2:
                self.filas_dobles += [r_et, r_val]
            for pos, cdef in fila:
                x = FICHA['x'][0] if pos in ('izq', 'completo') else FICHA['x'][1]
                w = FICHA['wFull'] if pos == 'completo' else FICHA['w']
                ci0, ci1 = self.col(x), self.col(x + FICHA['sepX']) - 1
                ct0, ct1 = self.col(x + FICHA['sepX']), self.col(x + w) - 1
                # zona del ícono (2 filas) con el separador a la derecha
                self.pintar(r_et, ci0, r_val, ct1, PAL['ficha'])
                sep = Side(style='thin', color='FF' + PAL['separador'])
                for r in (r_et, r_val):
                    cel = self.ws.cell(r, ci1)
                    cel.border = Border(right=sep)
                self.unir(r_et, ci0, r_val, ci1)
                lado = round(FICHA['iconSize'] * K)
                self.imagen('ficha_' + (cdef.get('icono') or 'file-text'), r_et, x + FICHA['iconCx'] - FICHA['iconSize'] / 2,
                            dy_px=(44 - FICHA['iconSize'] / 2) * K, ancho=lado, alto=lado)
                et = self.unir(r_et, ct0, r_et, ct1)
                et.value = cdef['etiqueta']
                et.font = Font(name=FUENTE, size=round(20.1 * K * 0.75 * 2) / 2, bold=True, color='FF' + PAL['azulEtiqueta'])
                et.alignment = Alignment(horizontal='left', vertical='bottom', indent=1)
                va = self.unir(r_val, ct0, r_val, ct1)
                va.value = cdef.get('valor') or None
                va.font = Font(name=FUENTE, size=round(25.2 * K * 0.75 * 2) / 2, bold=True, color='FF' + PAL['navy'])
                va.alignment = Alignment(horizontal='left', vertical='top', indent=1, wrap_text=True)
                celdas[cdef['etiqueta']] = f'${L(ct0)}${r_val}'
        # divisor vertical entre las dos columnas de fichas (como en el PDF)
        div = Side(style='thin', color='FF' + 'E6EEF5')
        c_div = self.col(709)
        for r in self.filas_dobles:
            cel = self.ws.cell(r, c_div)
            cel.border = Border(left=div)
        r_fin = self.nueva_fila(alto_u=17.5)
        # tarjeta blanca con borde
        for r in range(r_ini, r_fin + 1):
            for c in range(c_card0, c_card1 + 1):
                cel = self.ws.cell(r, c)
                if sin_relleno(cel):
                    cel.fill = fill(PAL['blanco'])
        self.borde_caja(r_ini, c_card0, r_fin, c_card1, PAL['borde'])
        return celdas

    def tabla(self, tabla, n_filas):
        K = self.K
        tam = {'sReg': 18.5, 'sBold': 19.5, 'sHead': 19.2, **(tabla.get('tam') or {})}
        cols = tabla['columnas']
        anchos = self._anchos_tabla()
        xs, acc = [], TABLA['x']
        for w in anchos:
            xs.append(acc); acc += w
        r_t = self.nueva_fila(alto_u=TABLA['hHead'])
        borde_col = Side(style='thin', color='FF2D6AA1')
        for i, c in enumerate(cols):
            c0, c1 = self.col(xs[i]), self.col(xs[i] + anchos[i]) - 1
            t0, t1 = (xs[i] - TABLA['x']) / TABLA['w'], (xs[i] + anchos[i] - TABLA['x']) / TABLA['w']
            g = GradientFill(type='linear', degree=0, stop=(Stop('FF' + color_en(t0), 0), Stop('FF' + color_en(t1), 1)))
            cel = self.unir(r_t, c0, r_t, c1)
            for cc in range(c0, c1 + 1):
                self.ws.cell(r_t, cc).fill = g
            cel.value = c['titulo']
            cel.font = Font(name=FUENTE, size=round(tam['sHead'] * K * 0.75 * 2) / 2, bold=True, color='FFFFFFFF')
            con_icono = bool(c.get('icono'))
            # Una sola palabra ("Aprobaciones"): se achica para caber en vez de partirse.
            una_palabra = ' ' not in c['titulo'].strip()
            cel.alignment = Alignment(horizontal='left' if con_icono else 'center', vertical='center',
                                      wrap_text=not una_palabra, shrink_to_fit=una_palabra,
                                      indent=max(1, round((63 - 4) * K / 9)) if con_icono else 0)
            if i:
                cel.border = Border(left=borde_col)
            if con_icono:
                lado = round(36 * K)
                self.imagen('tabla_' + c['icono'], r_t, xs[i] + 17, dy_px=(TABLA['hHead'] * K - lado) / 2, ancho=lado, alto=lado)
        sv = Side(style='thin', color='FF' + PAL['lineaTabla'])
        se = Side(style='thin', color='FF' + PAL['bordeTabla'])
        for k in range(n_filas):
            r = self.nueva_fila(alto_u=TABLA['hMin'])
            fondo = PAL['filaAlterna'] if k % 2 else PAL['blanco']
            for i, c in enumerate(cols):
                c0, c1 = self.col(xs[i]), self.col(xs[i] + anchos[i]) - 1
                cel = self.unir(r, c0, r, c1)
                for cc in range(c0, c1 + 1):
                    self.ws.cell(r, cc).fill = fill(fondo)
                if i == 0:
                    cel.value = k + 1
                negrita = c['estilo'] in ('fuerte', 'numero', 'pildora')
                color = PAL['numero'] if c['estilo'] == 'numero' else (PAL['navy'] if c['estilo'] in ('fuerte', 'pildora') else PAL['texto'])
                cel.font = Font(name=FUENTE, size=round((tam['sBold'] if negrita else tam['sReg']) * K * 0.75 * 2) / 2, bold=negrita, color='FF' + color)
                centro = c.get('alinear') == 'centro' or c['estilo'] == 'pildora'
                cel.alignment = Alignment(horizontal='center' if centro else 'left', vertical='center', wrap_text=True,
                                          indent=0 if centro else 1)
                cel.border = Border(left=se if i == 0 else sv, right=se if i == len(cols) - 1 else None,
                                    bottom=se if k == n_filas - 1 else None)
        return r_t

    def pie(self):
        K = self.K
        self.nueva_fila(alto_u=30)
        r_linea = self.nueva_fila(alto_u=12)
        self.imagen_estirada('pie', r_linea, r_linea)
        r = self.nueva_fila(alto_u=34)
        cel = self.ws.cell(r, self.col(38.5))
        doc = self.doc
        sz = round(18 * K * 0.75 * 2) / 2
        cel.value = CellRichText([
            TextBlock(InlineFont(rFont=FUENTE, sz=sz, b=True, color='FF' + PAL['pieFuerte']), 'METALIUM'),
            TextBlock(InlineFont(rFont=FUENTE, sz=sz, color='FF' + PAL['pieSep']), ' | '),
            TextBlock(InlineFont(rFont=FUENTE, sz=sz, color='FF' + PAL['pieTexto']), doc['pie']['codigo']),
            TextBlock(InlineFont(rFont=FUENTE, sz=sz, color='FF' + PAL['pieSep']), ' | '),
            TextBlock(InlineFont(rFont=FUENTE, sz=sz, color='FF' + PAL['pieTexto']), doc['pie']['nombre']),
        ])
        cel.alignment = Alignment(horizontal='left', vertical='center')
        self.fila_pie = r
        self.nueva_fila(alto_u=20)

    # ------------------------------------------------------------ armado
    def armar(self, aprobadores):
        self.encabezado()
        refs = {}
        for i, sec in enumerate(self.doc['secciones']):
            if i:
                self.nueva_fila(alto_u=30)
            self.barra(i + 1, sec['titulo'])
            if sec.get('fichas'):
                self.nueva_fila(alto_u=20)
                refs.update(self.fichas(sec['fichas']))
            else:
                self.nueva_fila(alto_u=21)
                self.r_titulos = self.tabla(sec['tabla'], aprobadores)
        self.cuadros_encabezado([refs['Código'], refs['Fecha del documento'], refs['Revisión']])
        self.pie()
        # fondo gris claro de la hoja donde no hay nada
        for r in range(1, self.fila):
            for c in range(1, self.col_fin + 1):
                cel = self.ws.cell(r, c)
                if sin_relleno(cel):
                    cel.fill = fill(PAL['fondo'])
        ws = self.ws
        ws.print_area = f'A1:{L(self.col_fin)}{self.fila - 1}'
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        if self.tipo == 'maestro':  # la fila de títulos de la tabla se repite en cada hoja
            ws.print_title_rows = f'{self.r_titulos}:{self.r_titulos}'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.25
        ws.page_margins.top = 0.3
        ws.page_margins.bottom = 0.45
        ws.oddFooter.right.text = 'Página &P de &N'
        ws.oddFooter.right.size = 7
        ws.oddFooter.right.font = FUENTE
        ws.oddFooter.right.color = PAL['pieTexto']


def xml_cuadro(cid, q):
    texto = (q['texto'] or '').replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    b = '1' if q['b'] else '0'
    rpr = (f'<a:rPr lang="es-CL" sz="{q["sz"]}" b="{b}" i="0" u="none" strike="noStrike">'
           f'<a:solidFill><a:srgbClr val="{q["color"]}"/></a:solidFill><a:latin typeface="{FUENTE}"/><a:cs typeface="{FUENTE}"/></a:rPr>')
    if q['vinculo']:
        guid = '{A1B2C3D4-0000-4000-8000-%012d}' % cid
        cuerpo = f'<a:fld id="{guid}" type="TxLink">{rpr}<a:pPr algn="{q["algn"]}"/><a:t>{texto}</a:t></a:fld>'
    else:
        cuerpo = f'<a:r>{rpr}<a:t>{texto}</a:t></a:r>'
    return (
        f'<xdr:oneCellAnchor><xdr:from><xdr:col>{q["col"]}</xdr:col><xdr:colOff>{q["colOff"]}</xdr:colOff>'
        f'<xdr:row>{q["row"]}</xdr:row><xdr:rowOff>{q["rowOff"]}</xdr:rowOff></xdr:from>'
        f'<xdr:ext cx="{q["cx"]}" cy="{q["cy"]}"/>'
        f'<xdr:sp macro="" textlink="{q["vinculo"] or ""}"><xdr:nvSpPr><xdr:cNvPr id="{cid}" name="{q["nombre"]}"/>'
        f'<xdr:cNvSpPr txBox="1"/></xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{q["cx"]}" cy="{q["cy"]}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></xdr:spPr>'
        f'<xdr:txBody><a:bodyPr vertOverflow="overflow" horzOverflow="overflow" wrap="{"square" if q["wrap"] else "none"}" '
        f'lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="ctr"/><a:lstStyle/>'
        f'<a:p><a:pPr algn="{q["algn"]}"/>{cuerpo}<a:endParaRPr lang="es-CL" sz="{q["sz"]}"/></a:p></xdr:txBody></xdr:sp>'
        f'<xdr:clientData/></xdr:oneCellAnchor>'
    )


def guardar(comp, ruta):
    tmp = ruta + '.tmp'
    comp.wb.save(tmp)
    zin = zipfile.ZipFile(tmp)
    zout = zipfile.ZipFile(ruta, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == 'xl/drawings/drawing1.xml':
            s = data.decode('utf-8')
            ns = ('xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                  'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" ')
            s = re.sub(r'<wsDr ', '<wsDr ' + ns, s, count=1)
            extra = ''.join(xml_cuadro(9001 + i, q) for i, q in enumerate(comp.cuadros))
            s = s.replace('</wsDr>', extra + '</wsDr>')
            data = s.encode('utf-8')
        zout.writestr(item, data)
    zout.close(); zin.close()
    os.remove(tmp)


# filas de la tabla: una por aprobador en cada comprobante; 20 filas en el maestro
APROBADORES = {'ingreso': 2, 'traslado': 3, 'aumento_sueldo': 1, 'cambio_cargo': 1, 'renovacion': 1, 'desvinculacion': 1,
               'cambio_turno': 1, 'maestro': 20}
ARCHIVO = {'ingreso': 'Ingreso', 'traslado': 'Traslado', 'aumento_sueldo': 'Aumento_de_sueldo', 'cambio_cargo': 'Cambio_de_cargo',
           'renovacion': 'Renovacion', 'desvinculacion': 'Desvinculacion', 'cambio_turno': 'Cambio_de_turno'}

if __name__ == '__main__':
    os.makedirs(os.path.join(BASE, 'out'), exist_ok=True)
    for tipo in sys.argv[1:]:
        comp = Comprobante(tipo)
        comp.armar(APROBADORES[tipo])
        nombre = (f'{comp.doc["encabezado"]["codigo"]}_Maestro_de_solicitudes.xlsx' if tipo == 'maestro'
                  else f'{comp.doc["encabezado"]["codigo"]}_Comprobante_{ARCHIVO[tipo]}.xlsx')
        guardar(comp, os.path.join(BASE, 'out', nombre))
        print(nombre)
